"""Document extractors for DOCX, Plain Text, Markdown, and HTML."""

import io
from pathlib import Path
import re
from typing import Any, Dict, List, Optional
from bs4 import BeautifulSoup
import docx

from ...exceptions import StorageError
from ...models.enums import SourceType
from .base import BaseExtractor
from .models import (
    ExtractedDocument,
    ExtractedHeading,
    ExtractedParagraph,
    ExtractedTable,
)


class DocxExtractor(BaseExtractor):
    """Structured extractor for Microsoft Word .docx documents using python-docx."""

    def can_handle(self, mime_type: str, filename: str) -> bool:
        lower_name = filename.lower()
        lower_mime = mime_type.lower()
        return lower_name.endswith(".docx") or "wordprocessingml.document" in lower_mime

    async def extract(
        self,
        source_id: str,
        filename: str,
        content: Optional[bytes] = None,
        metadata: Optional[Dict[str, Any]] = None,
        file_path: Optional[Path] = None,
    ) -> ExtractedDocument:
        try:
            if file_path and file_path.exists():
                doc = docx.Document(str(file_path))
            elif content is not None:
                doc = docx.Document(io.BytesIO(content))
            else:
                raise StorageError(f"Cannot extract DOCX '{filename}': no file content or path provided")
        except Exception as e:
            raise StorageError(f"Failed to parse DOCX document '{filename}': {e}") from e

        headings: List[ExtractedHeading] = []
        paragraphs: List[ExtractedParagraph] = []
        tables: List[ExtractedTable] = []
        raw_text_parts: List[str] = []

        current_section = None

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style and p.style.name else ""
            heading_match = re.match(r"^heading\s*([1-6])$", style_name)

            if heading_match:
                level = int(heading_match.group(1))
                headings.append(ExtractedHeading(text=text, level=level))
                current_section = text
                raw_text_parts.append(f"\n{'#' * level} {text}\n")
            else:
                paragraphs.append(ExtractedParagraph(text=text, section_title=current_section))
                raw_text_parts.append(text)

        # Extract Tables
        for table_idx, t in enumerate(doc.tables):
            if not t.rows:
                continue

            headers = [cell.text.strip() for cell in t.rows[0].cells]
            rows: List[List[str]] = []
            for row in t.rows[1:]:
                row_cells = [cell.text.strip() for cell in row.cells]
                rows.append(row_cells)

            table_name = f"Table {table_idx + 1}"
            tables.append(ExtractedTable(name=table_name, headers=headers, rows=rows))

            # Add table text representation
            raw_text_parts.append(f"\n[Table: {table_name}]")
            raw_text_parts.append(" | ".join(headers))
            for r in rows:
                raw_text_parts.append(" | ".join(r))
            raw_text_parts.append("")

        return ExtractedDocument(
            source_id=source_id,
            source_name=filename,
            source_type=SourceType.FILE,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headings=headings,
            paragraphs=paragraphs,
            tables=tables,
            raw_text="\n\n".join(raw_text_parts).strip(),
            metadata={"paragraph_count": len(paragraphs), "table_count": len(tables)},
        )


class TextMarkdownExtractor(BaseExtractor):
    """Structured extractor for Plain Text and Markdown files."""

    def can_handle(self, mime_type: str, filename: str) -> bool:
        lower_name = filename.lower()
        lower_mime = mime_type.lower()
        return (
            lower_name.endswith((".txt", ".md", ".markdown"))
            or "text/plain" in lower_mime
            or "text/markdown" in lower_mime
        )

    async def extract(
        self,
        source_id: str,
        filename: str,
        content: Optional[bytes] = None,
        metadata: Optional[Dict[str, Any]] = None,
        file_path: Optional[Path] = None,
    ) -> ExtractedDocument:
        if content is None:
            if not file_path or not file_path.exists():
                raise StorageError(f"Cannot extract text for '{filename}': no file content or path provided")
            content = file_path.read_bytes()

        # Decode UTF-8 with fallback
        try:
            text_str = content.decode("utf-8")
        except UnicodeDecodeError:
            text_str = content.decode("latin-1", errors="replace")

        lines = text_str.splitlines()
        headings: List[ExtractedHeading] = []
        paragraphs: List[ExtractedParagraph] = []
        tables: List[ExtractedTable] = []

        current_section: Optional[str] = None
        current_para_lines: List[str] = []

        def flush_paragraph():
            if current_para_lines:
                para_text = " ".join(current_para_lines).strip()
                if para_text:
                    paragraphs.append(ExtractedParagraph(text=para_text, section_title=current_section))
                current_para_lines.clear()

        in_table = False
        table_lines: List[str] = []

        for line in lines:
            trimmed = line.strip()

            # Markdown Table detection
            if trimmed.startswith("|") and trimmed.endswith("|"):
                flush_paragraph()
                in_table = True
                table_lines.append(trimmed)
                continue
            elif in_table:
                # End of table block
                in_table = False
                if len(table_lines) >= 2:
                    headers = [c.strip() for c in table_lines[0].strip("|").split("|")]
                    rows = []
                    for row_line in table_lines[1:]:
                        # Skip markdown divider row |---|---|
                        if re.match(r"^\|?\s*[-:]+[-| :]*\|?$", row_line):
                            continue
                        cells = [c.strip() for c in row_line.strip("|").split("|")]
                        rows.append(cells)
                    tables.append(ExtractedTable(name=f"Table {len(tables) + 1}", headers=headers, rows=rows))
                table_lines.clear()

            # Heading detection (# Heading)
            heading_match = re.match(r"^(#{1,6})\s+(.+)$", trimmed)
            if heading_match:
                flush_paragraph()
                level = len(heading_match.group(1))
                h_text = heading_match.group(2).strip()
                headings.append(ExtractedHeading(text=h_text, level=level))
                current_section = h_text
                continue

            if not trimmed:
                flush_paragraph()
            else:
                current_para_lines.append(trimmed)

        flush_paragraph()

        # Flush any trailing table
        if in_table and len(table_lines) >= 2:
            headers = [c.strip() for c in table_lines[0].strip("|").split("|")]
            rows = []
            for row_line in table_lines[1:]:
                if not re.match(r"^\|?\s*[-:]+[-| :]*\|?$", row_line):
                    rows.append([c.strip() for c in row_line.strip("|").split("|")])
            tables.append(ExtractedTable(name=f"Table {len(tables) + 1}", headers=headers, rows=rows))

        mime = "text/markdown" if filename.lower().endswith((".md", ".markdown")) else "text/plain"

        return ExtractedDocument(
            source_id=source_id,
            source_name=filename,
            source_type=SourceType.TEXT,
            mime_type=mime,
            headings=headings,
            paragraphs=paragraphs,
            tables=tables,
            raw_text=text_str.strip(),
            metadata={"line_count": len(lines), "heading_count": len(headings)},
        )


class HtmlExtractor(BaseExtractor):
    """Structured extractor for HTML documents and web pages using BeautifulSoup4."""

    def can_handle(self, mime_type: str, filename: str) -> bool:
        lower_name = filename.lower()
        lower_mime = mime_type.lower()
        return (
            lower_name.endswith((".html", ".htm", ".xhtml"))
            or "text/html" in lower_mime
            or "xhtml" in lower_mime
            or lower_name.startswith("http://")
            or lower_name.startswith("https://")
        )

    async def extract(
        self,
        source_id: str,
        filename: str,
        content: Optional[bytes] = None,
        metadata: Optional[Dict[str, Any]] = None,
        file_path: Optional[Path] = None,
    ) -> ExtractedDocument:
        if content is None:
            if not file_path or not file_path.exists():
                raise StorageError(f"Cannot extract HTML for '{filename}': no file content or path provided")
            content = file_path.read_bytes()

        soup = BeautifulSoup(content, "html.parser")

        # Strip non-content and noisy elements
        for element in soup(["script", "style", "nav", "footer", "header", "noscript", "svg", "form"]):
            element.decompose()

        headings: List[ExtractedHeading] = []
        paragraphs: List[ExtractedParagraph] = []
        tables: List[ExtractedTable] = []
        raw_text_parts: List[str] = []

        # Extract page title if available
        title_tag = soup.find("title")
        current_section = None
        if title_tag and title_tag.text.strip():
            page_title = title_tag.text.strip()
            headings.append(ExtractedHeading(text=page_title, level=1))
            current_section = page_title
            raw_text_parts.append(f"# {page_title}\n")

        # Traverse body elements
        body = soup.body or soup
        for el in body.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "blockquote", "table"]):
            tag_name = el.name.lower()
            if tag_name.startswith("h"):
                level = int(tag_name[1])
                text = el.get_text(separator=" ", strip=True)
                if text:
                    headings.append(ExtractedHeading(text=text, level=level))
                    current_section = text
                    raw_text_parts.append(f"\n{'#' * level} {text}\n")
            elif tag_name in ("p", "blockquote"):
                text = el.get_text(separator=" ", strip=True)
                if text:
                    paragraphs.append(ExtractedParagraph(text=text, section_title=current_section))
                    raw_text_parts.append(text)
            elif tag_name == "table":
                headers = []
                rows = []
                th_elements = el.find_all("th")
                if th_elements:
                    headers = [th.get_text(separator=" ", strip=True) for th in th_elements]

                for tr in el.find_all("tr"):
                    tds = tr.find_all("td")
                    if tds:
                        row_data = [td.get_text(separator=" ", strip=True) for td in tds]
                        rows.append(row_data)

                table_name = f"Table {len(tables) + 1}"
                tables.append(ExtractedTable(name=table_name, headers=headers, rows=rows))
                raw_text_parts.append(f"\n[Table: {table_name}]")
                if headers:
                    raw_text_parts.append(" | ".join(headers))
                for r in rows:
                    raw_text_parts.append(" | ".join(r))
                raw_text_parts.append("")

        is_url = filename.startswith("http://") or filename.startswith("https://") or "url" in metadata
        return ExtractedDocument(
            source_id=source_id,
            source_name=filename,
            source_type=SourceType.URL if is_url else SourceType.FILE,
            mime_type="text/html",
            headings=headings,
            paragraphs=paragraphs,
            tables=tables,
            raw_text="\n\n".join(raw_text_parts).strip(),
            metadata={"title": current_section, "headings_found": len(headings)},
        )
