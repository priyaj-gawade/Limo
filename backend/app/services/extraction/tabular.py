"""Tabular data extractors for Excel (.xlsx) and CSV/TSV spreadsheets."""

import csv
import io
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl

from ...exceptions import StorageError
from ...models.enums import SourceType
from .base import BaseExtractor
from .models import (
    ExtractedDocument,
    ExtractedHeading,
    ExtractedParagraph,
    ExtractedTable,
)

logger = logging.getLogger("limo.services.extraction.tabular")

# Cap on rows extracted per sheet to prevent memory exhaustion and token budget overrun
MAX_ROWS_PER_SHEET = 2000


class SpreadsheetExtractor(BaseExtractor):
    """Structured extractor for Excel (.xlsx) workbooks and CSV/TSV tables."""

    def can_handle(self, mime_type: str, filename: str) -> bool:
        lower_name = filename.lower()
        lower_mime = mime_type.lower()
        return (
            lower_name.endswith((".xlsx", ".csv", ".tsv"))
            or "spreadsheet" in lower_mime
            or "csv" in lower_mime
            or "tab-separated-values" in lower_mime
        )

    async def extract(
        self,
        source_id: str,
        filename: str,
        content: Optional[bytes] = None,
        metadata: Optional[Dict[str, Any]] = None,
        file_path: Optional[Path] = None,
    ) -> ExtractedDocument:
        lower_name = filename.lower()
        meta = metadata or {}

        if lower_name.endswith((".csv", ".tsv")) or "csv" in meta.get("mime_type", ""):
            return self._extract_delimited(source_id, filename, file_path, content)
        else:
            return self._extract_xlsx(source_id, filename, file_path, content)

    def _extract_delimited(
        self,
        source_id: str,
        filename: str,
        file_path: Optional[Path],
        content: Optional[bytes],
    ) -> ExtractedDocument:
        """Extract CSV or TSV file."""
        delimiter = "\t" if filename.lower().endswith(".tsv") else ","
        if file_path and file_path.exists():
            text_data = file_path.read_text(encoding="utf-8", errors="replace")
        elif content is not None:
            try:
                text_data = content.decode("utf-8")
            except UnicodeDecodeError:
                text_data = content.decode("latin-1", errors="replace")
        else:
            raise StorageError(f"Cannot extract delimited table for '{filename}': no file content or path provided")

        reader = csv.reader(io.StringIO(text_data), delimiter=delimiter)
        rows: List[List[str]] = []
        for r_idx, row in enumerate(reader):
            if r_idx >= MAX_ROWS_PER_SHEET:
                break
            rows.append([c.strip() for c in row])

        headers = rows[0] if rows else []
        data_rows = rows[1:] if len(rows) > 1 else []

        table = ExtractedTable(name=filename, headers=headers, rows=data_rows)
        raw_text_parts = [
            f"# Spreadsheet: {filename}",
            f"Columns: {', '.join(headers)}",
            f"Total Rows: {len(data_rows)}",
            "\nSample Data:",
            " | ".join(headers),
        ]
        for r in data_rows[:25]:  # Preview first 25 rows in raw text
            raw_text_parts.append(" | ".join(r))

        return ExtractedDocument(
            source_id=source_id,
            source_name=filename,
            source_type=SourceType.TABULAR,
            mime_type="text/csv",
            headings=[ExtractedHeading(text=filename, level=1)],
            tables=[table],
            raw_text="\n".join(raw_text_parts),
            metadata={"row_count": len(rows), "column_count": len(headers)},
        )

    def _extract_xlsx(
        self,
        source_id: str,
        filename: str,
        file_path: Optional[Path],
        content: Optional[bytes],
    ) -> ExtractedDocument:
        """Extract multi-sheet Excel workbook using openpyxl."""
        if file_path and file_path.exists():
            workbook = openpyxl.load_workbook(filename=str(file_path), data_only=True, read_only=True)
        elif content is not None:
            workbook = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True, read_only=True)
        else:
            raise StorageError(f"Cannot extract Excel workbook '{filename}': no file content or path provided")

        headings: List[ExtractedHeading] = []
        paragraphs: List[ExtractedParagraph] = []
        tables: List[ExtractedTable] = []
        raw_text_parts: List[str] = [f"# Workbook: {filename}"]

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            headings.append(ExtractedHeading(text=f"Sheet: {sheet_name}", level=2))
            raw_text_parts.append(f"\n## Sheet: {sheet_name}")

            sheet_rows: List[List[str]] = []
            for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
                if r_idx >= MAX_ROWS_PER_SHEET:
                    break
                row_str = [str(c).strip() if c is not None else "" for c in row]
                # Filter out completely empty rows
                if any(row_str):
                    sheet_rows.append(row_str)

            if not sheet_rows:
                paragraphs.append(ExtractedParagraph(text="[Empty Sheet]", section_title=sheet_name))
                continue

            headers = sheet_rows[0]
            data_rows = sheet_rows[1:] if len(sheet_rows) > 1 else []

            tables.append(ExtractedTable(name=sheet_name, headers=headers, rows=data_rows))

            raw_text_parts.append(f"Headers: {', '.join(headers)}")
            raw_text_parts.append(f"Rows: {len(data_rows)}")
            raw_text_parts.append(" | ".join(headers))
            for r in data_rows[:20]:
                raw_text_parts.append(" | ".join(r))

        return ExtractedDocument(
            source_id=source_id,
            source_name=filename,
            source_type=SourceType.TABULAR,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headings=headings,
            paragraphs=paragraphs,
            tables=tables,
            raw_text="\n".join(raw_text_parts),
            metadata={"sheet_count": len(workbook.sheetnames), "table_count": len(tables)},
        )
